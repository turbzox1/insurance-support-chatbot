from pathlib import Path

import pytest
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings

from src.retrieval.context_compressor import ContextCompressor
from src.retrieval.document_loader import load_documents
from src.retrieval.hybrid_retriever import HybridRetriever
from src.retrieval.ingest import ingest
from src.retrieval.reranker import Reranker


class FakeEmbeddings(Embeddings):
    def embed_documents(self, texts):
        return [self.embed_query(text) for text in texts]

    def embed_query(self, text):
        words = text.lower()
        return [
            float(words.count("insurance")),
            float(words.count("claim")),
            float(words.count("copay")),
            0.1,
        ]


def test_load_supplied_documents():
    docs = load_documents()
    assert len({doc.metadata["source"] for doc in docs}) == 5
    zero = next(doc for doc in docs if doc.metadata["source"].endswith("Zero Co-Pay.pdf"))
    assert zero.metadata["page"] == 0 and "cannot be bought" in zero.page_content


def test_optional_folders_and_docx(tmp_path):
    from docx import Document as WordDocument

    assert load_documents(tmp_path) == []
    word = WordDocument()
    word.add_paragraph("Insurance coverage")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "Claim condition"
    word.save(tmp_path / "example.docx")
    docs = load_documents(tmp_path)
    assert "Claim condition" in docs[0].page_content
    assert docs[0].metadata["source"] == "example.docx"


def test_real_chroma_roundtrip_idempotence_and_hybrid(tmp_path):
    from langchain_chroma import Chroma

    data = tmp_path / "data"
    data.mkdir()
    (data / "policy.txt").write_text(
        "Insurance copay benefit. Claim documents required.", encoding="utf-8"
    )
    (data / "other.txt").write_text("Insurance premium renewal terms.", encoding="utf-8")
    store = Chroma(
        collection_name="test_insurance",
        persist_directory=str(tmp_path / "db"),
        embedding_function=FakeEmbeddings(),
    )
    stats = ingest(data, store)
    ids = store.get()["ids"]
    ingest(data, store)
    assert set(store.get()["ids"]) == set(ids) and stats["chunks"] == 2
    hybrid = HybridRetriever(store)
    assert hybrid.bm25_search("copay")[0][0].metadata["source"] == "policy.txt"
    results = hybrid.hybrid_search("copay claim", k=2)
    assert len(results) == 2 and results[0].metadata["source"] == "policy.txt"
    assert len({doc.metadata["chunk_id"] for doc in results}) == len(results)
    # Refresh removes only stale indexed chunks after new data is safely upserted.
    (data / "policy.txt").write_text("Insurance claim updated.", encoding="utf-8")
    ingest(data, store)
    assert len(store.get()["ids"]) == 2
    assert any("updated" in text for text in store.get()["documents"])
    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(ValueError):
        ingest(empty, store)
    assert len(store.get()["ids"]) == 2


def test_hybrid_empty_index():
    class Store:
        def get(self):
            return {"ids": [], "documents": [], "metadatas": []}

    assert HybridRetriever(Store()).hybrid_search("claim") == []


def test_rrf_preserves_same_text_different_sources():
    class Store:
        def get(self):
            return {
                "ids": ["a", "b"],
                "documents": ["same insurance", "same insurance"],
                "metadatas": [{"source": "a"}, {"source": "b"}],
            }

        def similarity_search_with_score(self, query, k):
            return [
                (
                    Document(
                        page_content="same insurance", metadata={"source": "b", "chunk_id": "b"}
                    ),
                    0.1,
                )
            ]

    results = HybridRetriever(Store()).hybrid_search("insurance")
    assert len(results) == 2 and results[0].metadata["source"] == "b"


def test_reranking_empty_and_order():
    class Model:
        def predict(self, pairs):
            return [-1, 5]

    ranker = Reranker(Model())
    assert ranker.rerank("q", []) == []
    docs = [Document(page_content="a"), Document(page_content="b")]
    assert ranker.rerank("q", docs)[0] == (docs[1], 5.0)


def test_compression_deduplicates_and_preserves_metadata():
    docs = [
        Document(page_content=text, metadata={"source": str(i), "page": i})
        for i, text in enumerate(["insurance insurance", "insurance insurance", "claim claim"])
    ]
    compressed = ContextCompressor(FakeEmbeddings()).compress(docs, max_chars=25)
    assert len(compressed) == 2
    assert compressed[1].metadata["page"] == 2
    assert sum(len(doc.page_content) for doc in compressed) <= 25
    assert docs[2].page_content == "claim claim"


def test_paths_are_independent_of_cwd(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    assert len({doc.metadata["source"] for doc in load_documents()}) == 5
    from src.config.config import VECTORSTORE_PATH

    assert Path(VECTORSTORE_PATH).is_absolute()

"""Load PDF, UTF-8 text and DOCX files with source/page metadata."""

from pathlib import Path

from langchain_core.documents import Document

from src.config.config import BASE_DIR


def load_documents(data_dir=None):
    root = Path(data_dir) if data_dir else BASE_DIR / "data"
    documents = []
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        source = path.relative_to(root).as_posix()
        if path.suffix.lower() == ".pdf":
            from pypdf import PdfReader

            for page, item in enumerate(PdfReader(path).pages):
                text = item.extract_text() or ""
                if text.strip():
                    documents.append(
                        Document(page_content=text, metadata={"source": source, "page": page})
                    )
        elif path.suffix.lower() == ".txt":
            text = path.read_text(encoding="utf-8")
            if text.strip():
                documents.append(Document(page_content=text, metadata={"source": source}))
        elif path.suffix.lower() == ".docx":
            from docx import Document as WordDocument

            word = WordDocument(path)
            text = "\n".join(
                [p.text for p in word.paragraphs]
                + [
                    " | ".join(c.text for c in row.cells)
                    for table in word.tables
                    for row in table.rows
                ]
            )
            if text.strip():
                documents.append(Document(page_content=text, metadata={"source": source}))
    return documents
